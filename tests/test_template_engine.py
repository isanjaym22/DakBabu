"""Tests for src.core.template_engine.

Covers placeholder scanning, case-insensitive matching, multiple occurrence
replacement, cross-validation (missing and unused columns), and the
python-docx split-runs edge case.
"""

import pytest
from pathlib import Path

from docx import Document

from src.core.template_engine import (
    scan_placeholders,
    cross_validate_placeholders,
    replace_placeholders,
    scan_docx_placeholders,
    generate_personalized_docx,
)
from src.core.models import Recipient, RecipientStatus


@pytest.fixture
def tmp_docx(tmp_path: Path):
    """Factory fixture that creates a .docx file from paragraph strings.

    Usage:
        path = tmp_docx(["Dear {{Name}},", "From {{Department}}."])
    """

    def _create(
        paragraphs: list[str],
        filename: str = "template.docx",
    ) -> Path:
        """Write paragraphs to a temporary .docx file.

        Args:
            paragraphs: List of paragraph text strings.
            filename: Name for the created file.

        Returns:
            Path to the created .docx file.
        """
        doc = Document()
        for text in paragraphs:
            doc.add_paragraph(text)
        file_path = tmp_path / filename
        doc.save(str(file_path))
        return file_path

    return _create


@pytest.fixture
def tmp_docx_split_runs(tmp_path: Path):
    """Factory fixture that creates a .docx with a placeholder split across runs.

    Simulates the common python-docx problem where Word splits text
    mid-placeholder due to formatting changes.
    """

    def _create() -> Path:
        """Create a .docx where {{name}} is split across 3 runs.

        Returns:
            Path to the created .docx file.
        """
        doc = Document()
        paragraph = doc.add_paragraph()

        # Simulate "Hello {{name}}, welcome!" split across runs:
        # Run 1: "Hello {{"
        # Run 2: "name"
        # Run 3: "}}, welcome!"
        run1 = paragraph.add_run("Hello {{")
        run2 = paragraph.add_run("name")
        run2.bold = True  # Different formatting causes the split
        run3 = paragraph.add_run("}}, welcome!")

        file_path = tmp_path / "split_runs.docx"
        doc.save(str(file_path))
        return file_path

    return _create


class TestScanPlaceholders:
    """Tests for the scan_placeholders function."""

    def test_finds_all_placeholders(self) -> None:
        """Scans and returns all placeholder names from text."""
        text = "Dear {{Name}}, you are in {{Department}} with roll {{Roll_No}}."
        result = scan_placeholders(text)

        assert result == {"name", "department", "roll_no"}

    def test_case_insensitive_matching(self) -> None:
        """Placeholders are returned lowercased regardless of input case."""
        text = "{{NAME}} and {{name}} and {{Name}}"
        result = scan_placeholders(text)

        assert result == {"name"}

    def test_multiple_occurrences_counted_once(self) -> None:
        """Same placeholder appearing multiple times is in the set once."""
        text = "{{name}} appears and {{name}} appears again."
        result = scan_placeholders(text)

        assert result == {"name"}
        assert len(result) == 1

    def test_no_placeholders_returns_empty(self) -> None:
        """Text without placeholders returns an empty set."""
        text = "No placeholders here."
        result = scan_placeholders(text)

        assert result == set()

    def test_empty_string_returns_empty(self) -> None:
        """Empty string returns an empty set."""
        assert scan_placeholders("") == set()

    def test_strips_whitespace_in_placeholder(self) -> None:
        """Whitespace inside braces is stripped from the name."""
        text = "{{ name }} and {{  department  }}"
        result = scan_placeholders(text)

        assert result == {"name", "department"}


class TestCrossValidatePlaceholders:
    """Tests for the cross_validate_placeholders function."""

    def test_all_placeholders_have_columns(self) -> None:
        """No errors when all placeholders match columns."""
        placeholders = {"name", "department"}
        columns = {"name", "email", "department"}

        result = cross_validate_placeholders(placeholders, columns)

        assert result.is_valid is True
        assert len(result.errors) == 0

    def test_missing_columns_produce_errors(self) -> None:
        """Placeholders without matching columns produce clear errors."""
        placeholders = {"name", "department", "roll_no"}
        columns = {"name", "email"}

        result = cross_validate_placeholders(placeholders, columns)

        assert result.is_valid is False
        assert len(result.errors) == 2
        assert any("department" in e.lower() for e in result.errors)
        assert any("roll_no" in e.lower() for e in result.errors)

    def test_unused_columns_no_error(self) -> None:
        """Columns not used as placeholders do not cause errors."""
        placeholders = {"name"}
        columns = {"name", "email", "department", "roll_no"}

        result = cross_validate_placeholders(placeholders, columns)

        assert result.is_valid is True
        assert len(result.errors) == 0

    def test_email_column_excluded_from_placeholders(self) -> None:
        """The 'email' column is not available as a placeholder."""
        placeholders = {"email"}
        columns = {"email", "name"}

        result = cross_validate_placeholders(placeholders, columns)

        assert result.is_valid is False
        assert any("email" in e.lower() for e in result.errors)

    def test_empty_placeholders_always_valid(self) -> None:
        """No placeholders means nothing to validate — always valid."""
        result = cross_validate_placeholders(set(), {"name", "email"})

        assert result.is_valid is True


class TestReplacePlaceholders:
    """Tests for the replace_placeholders function."""

    def test_replaces_all_occurrences(self) -> None:
        """All occurrences of the same placeholder are replaced."""
        text = "Dear {{name}}, hello {{name}}! Welcome, {{name}}."
        data = {"name": "Alice"}

        result = replace_placeholders(text, data)

        assert result == "Dear Alice, hello Alice! Welcome, Alice."
        assert "{{" not in result

    def test_case_insensitive_replacement(self) -> None:
        """Replacement is case-insensitive."""
        text = "{{Name}} and {{NAME}} and {{name}}"
        data = {"name": "Bob"}

        result = replace_placeholders(text, data)

        assert result == "Bob and Bob and Bob"

    def test_multiple_different_placeholders(self) -> None:
        """Multiple different placeholders are each replaced correctly."""
        text = "{{name}} from {{department}}, roll {{roll_no}}"
        data = {"name": "Alice", "department": "CS", "roll_no": "42"}

        result = replace_placeholders(text, data)

        assert result == "Alice from CS, roll 42"

    def test_unmatched_placeholders_left_as_is(self) -> None:
        """Placeholders without matching data are left unchanged."""
        text = "Hello {{name}}, your code is {{code}}"
        data = {"name": "Alice"}

        result = replace_placeholders(text, data)

        assert result == "Hello Alice, your code is {{code}}"


class TestScanDocxPlaceholders:
    """Tests for the scan_docx_placeholders function."""

    def test_finds_placeholders_in_paragraphs(self, tmp_docx) -> None:
        """Scans paragraphs and finds all placeholders."""
        path = tmp_docx([
            "Dear {{Name}},",
            "Welcome to {{Department}}.",
            "Your roll number is {{Roll_No}}.",
        ])

        placeholders, error = scan_docx_placeholders(path)

        assert error is None
        assert placeholders == {"name", "department", "roll_no"}

    def test_handles_split_runs(self, tmp_docx_split_runs) -> None:
        """Detects placeholder that is split across multiple Word runs."""
        path = tmp_docx_split_runs()

        placeholders, error = scan_docx_placeholders(path)

        assert error is None
        assert "name" in placeholders

    def test_invalid_file_returns_error(self, tmp_path) -> None:
        """Non-docx file produces a clear error."""
        fake_path = tmp_path / "not_a_doc.docx"
        fake_path.write_text("This is not a docx file")

        placeholders, error = scan_docx_placeholders(fake_path)

        assert len(placeholders) == 0
        assert error is not None

    def test_file_not_found_returns_error(self) -> None:
        """Missing file produces a clear error."""
        placeholders, error = scan_docx_placeholders(
            Path("/nonexistent/doc.docx")
        )

        assert len(placeholders) == 0
        assert error is not None

    def test_finds_placeholders_in_tables(self, tmp_path) -> None:
        """Scans table cells and finds placeholders inside them."""
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].paragraphs[0].text = "Name: {{name}}"
        table.rows[0].cells[1].paragraphs[0].text = "Dept: {{department}}"
        file_path = tmp_path / "table_doc.docx"
        doc.save(str(file_path))

        placeholders, error = scan_docx_placeholders(file_path)

        assert error is None
        assert placeholders == {"name", "department"}


class TestGeneratePersonalizedDocx:
    """Tests for the generate_personalized_docx function."""

    def test_replaces_placeholders_in_output(self, tmp_docx, tmp_path) -> None:
        """Generated doc has all placeholders replaced with data."""
        template = tmp_docx([
            "Dear {{Name}},",
            "Welcome to {{Department}}.",
        ])
        output = tmp_path / "output" / "personalized.docx"
        recipient = Recipient(
            name="Alice",
            email="alice@example.com",
            extra_fields={"department": "Computer Science"},
        )

        success, error = generate_personalized_docx(
            template, recipient, output
        )

        assert success is True
        assert error == ""
        assert output.exists()

        # Verify the content
        doc = Document(str(output))
        texts = [p.text for p in doc.paragraphs]
        assert "Dear Alice," in texts
        assert "Welcome to Computer Science." in texts

    def test_multiple_occurrences_all_replaced(self, tmp_docx, tmp_path) -> None:
        """All occurrences of the same placeholder are replaced."""
        template = tmp_docx([
            "{{name}} is here. Hello {{name}}! Goodbye {{name}}.",
        ])
        output = tmp_path / "output.docx"
        recipient = Recipient(name="Bob", email="bob@example.com")

        success, error = generate_personalized_docx(
            template, recipient, output
        )

        assert success is True
        doc = Document(str(output))
        text = doc.paragraphs[0].text
        assert text == "Bob is here. Hello Bob! Goodbye Bob."

    def test_split_runs_handled(self, tmp_docx_split_runs, tmp_path) -> None:
        """Placeholders split across runs are still replaced."""
        template = tmp_docx_split_runs()
        output = tmp_path / "output.docx"
        recipient = Recipient(name="Charlie", email="c@d.com")

        success, error = generate_personalized_docx(
            template, recipient, output
        )

        assert success is True
        doc = Document(str(output))
        full_text = "".join(r.text for r in doc.paragraphs[0].runs)
        assert "Charlie" in full_text
        assert "{{name}}" not in full_text

    def test_invalid_template_returns_error(self, tmp_path) -> None:
        """Invalid template file returns failure without crashing."""
        fake = tmp_path / "bad.docx"
        fake.write_text("not a docx")
        output = tmp_path / "out.docx"
        recipient = Recipient(name="X", email="x@y.com")

        success, error = generate_personalized_docx(
            fake, recipient, output
        )

        assert success is False
        assert len(error) > 0
